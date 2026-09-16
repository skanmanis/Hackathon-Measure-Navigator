from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Iterable

from config import Settings
from llm import LLMAdapter
from models import Chunk
from retrieval import KnowledgeStore


AUTHORITIES = {"specification": 1, "vsd": 2, "mld": 2, "sql": 3, "sql_explanation": 4, "faq": 5}


class IngestionService:
    def __init__(self, settings: Settings, llm: LLMAdapter, store: KnowledgeStore):
        self.settings = settings
        self.llm = llm
        self.store = store

    def build(self) -> dict:
        chunks: list[Chunk] = []
        errors: list[dict] = []
        for measure in self.settings.measures:
            (self.settings.knowledge_root / measure).mkdir(parents=True, exist_ok=True)
        files = [path for path in self.settings.knowledge_root.rglob("*") if path.is_file()]
        for path in files:
            try:
                chunks.extend(self._chunks_for(path))
            except Exception as exc:
                errors.append({"file": str(path), "error": str(exc)})
        chunks.extend(self._faq_chunks())
        self.store.clear()
        if chunks:
            self.store.upsert(chunks)
        report = {
            "status": "ready" if chunks else "empty",
            "backend": self.store.backend,
            "files_seen": len(files),
            "chunks_indexed": len(chunks),
            "errors": errors,
            "llm_mode": self.llm.mode,
        }
        (self.settings.root / "data" / "index-report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
        return report

    def status(self) -> dict:
        report_path = self.settings.root / "data" / "index-report.json"
        report = json.loads(report_path.read_text(encoding="utf-8")) if report_path.exists() else {}
        report.update({"backend": self.store.backend, "chunks_indexed": self.store.count(), "llm_mode": self.llm.mode})
        report.setdefault("status", "ready" if self.store.count() else "not_indexed")
        return report

    def _chunks_for(self, path: Path) -> list[Chunk]:
        measure = self._measure_for(path)
        if not measure:
            raise ValueError("File must be placed under a CBP, GSD, KED, EED, or SPC folder")
        suffix = path.suffix.lower()
        if suffix == ".sql":
            return self._sql_chunks(path, measure)
        text, source_type = self._read_document(path)
        return self._text_chunks(path, measure, source_type, text)

    def _measure_for(self, path: Path) -> str | None:
        upper_parts = {part.upper() for part in path.parts}
        return next((measure for measure in self.settings.measures if measure in upper_parts), None)

    @staticmethod
    def _source_type(path: Path) -> str:
        name = path.name.lower()
        if "value" in name or "vsd" in name:
            return "vsd"
        if "medication" in name or "mld" in name:
            return "mld"
        if path.suffix.lower() == ".sql":
            return "sql"
        return "specification"

    def _read_document(self, path: Path) -> tuple[str, str]:
        suffix = path.suffix.lower()
        source_type = self._source_type(path)
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="ignore"), source_type
        if suffix == ".pdf":
            from pypdf import PdfReader

            pages = []
            for number, page in enumerate(PdfReader(str(path)).pages, 1):
                pages.append(f"[[PAGE {number}]]\n{page.extract_text() or ''}")
            return "\n\n".join(pages), source_type
        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                blocks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            return "\n\n".join(blocks), source_type
        if suffix == ".csv":
            with path.open(encoding="utf-8-sig", errors="ignore", newline="") as handle:
                return "\n".join(" | ".join(row) for row in csv.reader(handle)), source_type
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            blocks = []
            for sheet in workbook.worksheets:
                blocks.append(f"[[SHEET {sheet.title}]]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value not in (None, "")]
                    if values:
                        blocks.append(" | ".join(values))
            return "\n".join(blocks), source_type
        raise ValueError(f"Unsupported extension: {suffix}")

    def _text_chunks(self, path: Path, measure: str, source_type: str, text: str) -> list[Chunk]:
        sections = self._split_text(text)
        return [
            Chunk(
                chunk_id=self._id(path, index, source_type), text=section["text"], measure_id=measure,
                measurement_year=2026, source_type=source_type, source_name=path.name,
                locator=section["locator"], authority=AUTHORITIES[source_type], title=section.get("title", "")
            )
            for index, section in enumerate(sections, 1)
        ]

    @staticmethod
    def _split_text(text: str, target: int = 1400, overlap: int = 180) -> list[dict]:
        lines = [line.strip() for line in text.splitlines()]
        chunks: list[dict] = []
        buffer: list[str] = []
        locator = "Document"
        title = ""
        for line in lines:
            marker = re.match(r"\[\[(PAGE|SHEET) (.+?)\]\]", line)
            if marker:
                locator = f"{marker.group(1).title()} {marker.group(2)}"
                continue
            if re.match(r"^(#{1,4}\s+|\d+(?:\.\d+)*\s+|[A-Z][A-Z\s]{5,})", line) and len(line) < 160:
                title = line.lstrip("# ")
            if not line:
                continue
            buffer.append(line)
            joined = "\n".join(buffer)
            if len(joined) >= target:
                chunks.append({"text": joined, "locator": locator, "title": title})
                tail = joined[-overlap:]
                buffer = [tail]
        if buffer:
            chunks.append({"text": "\n".join(buffer), "locator": locator, "title": title})
        return chunks

    def _sql_chunks(self, path: Path, measure: str) -> list[Chunk]:
        sql = path.read_text(encoding="utf-8", errors="ignore")
        dependencies = sorted(set(re.findall(r"\b(?:EXEC(?:UTE)?|FROM|JOIN)\s+(?:\[?\w+\]?\.){0,2}\[?([A-Za-z_]\w+)\]?", sql, re.I)))
        blocks = self._split_sql(sql)
        chunks: list[Chunk] = []
        for index, (start_line, end_line, block) in enumerate(blocks, 1):
            fallback = self._fallback_sql_explanation(block)
            result = self.llm.structured(
                "Translate only the supplied T-SQL block into precise plain English. Do not add HEDIS rules absent from the SQL. Return JSON with explanation, direct_behavior, inherited_behavior, uncertainty.",
                block,
                {"explanation": fallback, "direct_behavior": fallback, "inherited_behavior": "", "uncertainty": ""},
            )
            explanation = result.get("explanation") or fallback
            lineage = f"T-SQL lines {start_line}-{end_line}"
            combined = f"English explanation: {explanation}\n\nOriginal T-SQL:\n{block}"
            chunks.append(Chunk(
                chunk_id=self._id(path, index, "sql_explanation"), text=combined, measure_id=measure,
                measurement_year=2026, source_type="sql_explanation", source_name=path.name,
                locator=lineage, authority=AUTHORITIES["sql_explanation"], title=f"SQL logical block {index}",
                sql_text=block, dependency_names=dependencies
            ))
        return chunks

    @staticmethod
    def _split_sql(sql: str) -> list[tuple[int, int, str]]:
        lines = sql.splitlines()
        starts = [0]
        pattern = re.compile(r"^\s*(?:WITH\s+\w+\s+AS|SELECT\b|INSERT\b|UPDATE\b|DELETE\b|MERGE\b|IF\b|BEGIN\b|EXEC(?:UTE)?\b)", re.I)
        for index, line in enumerate(lines):
            if index and pattern.search(line) and index - starts[-1] >= 5:
                starts.append(index)
        starts.append(len(lines))
        return [(starts[i] + 1, starts[i + 1], "\n".join(lines[starts[i] : starts[i + 1]]).strip()) for i in range(len(starts) - 1) if "\n".join(lines[starts[i] : starts[i + 1]]).strip()]

    @staticmethod
    def _fallback_sql_explanation(block: str) -> str:
        actions = []
        upper = block.upper()
        if "JOIN" in upper:
            actions.append("joins related data sources")
        if "WHERE" in upper:
            actions.append("filters records using the stated conditions")
        if "CASE" in upper:
            actions.append("derives a value through conditional logic")
        if any(term in upper for term in ("ROW_NUMBER", "RANK(", "MIN(", "MAX(")):
            actions.append("ranks or aggregates candidate records")
        if "EXEC" in upper:
            actions.append("invokes another stored procedure")
        return "This block " + ", then ".join(actions) + "." if actions else "This block contains T-SQL implementation logic that requires source-level review."

    def _faq_chunks(self) -> list[Chunk]:
        chunks = []
        for index, row in enumerate(self.store.list_faqs(), 1):
            raw_measure = row.get("measure_id", "CBP").replace("SYN-", "")
            if raw_measure not in self.settings.measures:
                continue
            text = f"Question: {row.get('question', '')}\nAnswer: {row.get('answer', '')}"
            chunks.append(Chunk(
                chunk_id=f"faq-{index}", text=text, measure_id=raw_measure, measurement_year=2026,
                source_type="faq", source_name="Hackathon FAQ", locator=row.get("source_refs", "FAQ"),
                authority=AUTHORITIES["faq"], title=row.get("question", "")
            ))
        return chunks

    @staticmethod
    def _id(path: Path, index: int, source_type: str) -> str:
        digest = hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:10]
        return f"{source_type}-{digest}-{index:04d}"
